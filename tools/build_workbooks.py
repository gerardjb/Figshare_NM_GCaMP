"""Rebuild Source Data workbooks from the packaged CSV files.

Tracked workbooks provide the sheet layout and formatting. Registered data
blocks are populated from CSVs, including the declared time-series reductions.
Only changed cell values are written; other workbook contents are preserved.
"""

from __future__ import annotations

import argparse
import csv
import html
import json
import math
import re
import shutil
import sys
import tempfile
import zipfile
from collections import OrderedDict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET

import numpy as np
from openpyxl import load_workbook

from publication_config import ITEM_BY_SOURCE_DATA, REPO_ROOT
from source_data_common import load_registry, local_csv_path, sha256


TIME_COLUMNS = {
    "time", "time_s", "time_ms", "stored_time", "simulation_relative_time",
    "plot_time_s", "time_from_last_stimulus_ms", "time_from_first_stimulus_ms", "x",
}
INDEX_COLUMNS = {"source_row_index", "source_time_index", "index", "source_point_index"}
GROUP_COLUMNS = {
    "axis_index", "series_index", "protocol_index", "state_group", "segment",
    "curve", "indicator", "calcium", "calcium_nm", "calcium_condition_nm",
    "saved_calcium_nm", "source_trace_index", "source_event_row",
    "structure_source_index", "stimulus_number", "display_snr", "spike_bin",
    "n_events", "n_all_events_before_plot_snr_filter", "method_key", "method",
    "spike_bin_label", "cell_type", "trace_kind", "trace_type", "condition",
    "cluster", "kinetics", "color_role", "saved_structure_type", "panel",
    "sampling_phase", "source_path", "source_sha256", "draw_id",
}
UNIT_SUFFIX = re.compile(r"\s*(?:\[[^]]+\]|\([^)]*\))\s*$")


@dataclass
class Block:
    rows: list[dict[str, Any]]
    mode: str
    dataset_label: str | None = None


def report_path(path: Path) -> str:
    """Use repository-relative paths when possible, absolute paths otherwise."""
    resolved = path.resolve()
    try:
        return resolved.relative_to(REPO_ROOT.resolve()).as_posix()
    except ValueError:
        return resolved.as_posix()


def parse(value: str) -> Any:
    stripped = value.strip()
    if stripped == "" or stripped.lower() in {"nan", "na", "n.a."}:
        return None
    if stripped.lower() == "true":
        return True
    if stripped.lower() == "false":
        return False
    try:
        number = float(stripped)
    except ValueError:
        return value
    if math.isfinite(number) and number.is_integer() and not any(
        ch in stripped.lower() for ch in (".", "e")
    ):
        return int(number)
    return number


def normalized_name(value: Any) -> str:
    text = UNIT_SUFFIX.sub("", str(value or "")).strip()
    text = text.replace("Δ", "Delta").replace("δ", "delta")
    return re.sub(r"[^0-9A-Za-z]+", "_", text).strip("_").lower()


def read_csv(
    dataset: dict[str, Any], csv_overrides: dict[str, Path] | None = None
) -> tuple[list[str], list[dict[str, Any]]]:
    path = (csv_overrides or {}).get(str(dataset["csv"]), local_csv_path(dataset))
    manifest_header = [str(entry["name"]) for entry in dataset["columns"]]
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        raw_header = next(reader)
        if len(raw_header) != len(manifest_header):
            raise ValueError(f"{dataset['csv']}: CSV/manifest column count differs")
        for raw, expected in zip(raw_header, manifest_header):
            if normalized_name(raw) != normalized_name(expected):
                raise ValueError(
                    f"{dataset['csv']}: column {raw!r} does not match manifest {expected!r}"
                )
        rows = [
            {name: parse(value) for name, value in zip(manifest_header, raw_row)}
            for raw_row in reader
        ]
    if len(rows) != int(dataset["rows_count"]):
        raise ValueError(f"{dataset['csv']}: row count differs from manifest")
    return manifest_header, rows


def is_number(value: Any) -> bool:
    return isinstance(value, (int, float, np.number)) and not isinstance(value, bool)


def resolve_sheet(workbook: Any, requested: str) -> Any:
    if requested in workbook.sheetnames:
        return workbook[requested]
    key = re.sub(r"[^0-9A-Za-z]+", "-", requested).strip("-")
    matches = [
        name for name in workbook.sheetnames
        if re.sub(r"[^0-9A-Za-z]+", "-", name).strip("-") == key
    ]
    if len(matches) != 1:
        raise KeyError(f"Could not resolve worksheet {requested!r}; found {workbook.sheetnames}")
    return workbook[matches[0]]


def unique_time_rows(
    rows: list[dict[str, Any]], time_column: str, columns: Iterable[str]
) -> list[dict[str, Any]]:
    ordered = sorted(rows, key=lambda row: float(row[time_column]))
    result: list[dict[str, Any]] = []
    for row in ordered:
        if result and float(row[time_column]) == float(result[-1][time_column]):
            if any(row.get(column) != result[-1].get(column) for column in columns):
                raise ValueError(f"Nonidentical duplicate timestamp {row[time_column]}")
            continue
        result.append(row)
    return result


def resample_segment(
    rows: list[dict[str, Any]],
    *,
    columns: list[str],
    time_column: str,
    group_columns: list[str],
    count: int,
    bounds: tuple[float, float] | None = None,
) -> list[dict[str, Any]]:
    if len(rows) == count and bounds is None:
        return [{column: row.get(column) for column in columns} for row in rows]
    if any(not is_number(row.get(time_column)) for row in rows):
        raise ValueError(f"Missing/non-numeric time in {time_column}")
    source = unique_time_rows(rows, time_column, columns)
    x = np.asarray([float(row[time_column]) for row in source], dtype=float)
    q = np.linspace(*(bounds or (x[0], x[-1])), count)
    if q[0] < x[0] or q[-1] > x[-1]:
        raise ValueError(f"Requested time bounds {bounds} exceed CSV range {(x[0], x[-1])}")
    numeric_columns = [
        column for column in columns
        if column not in group_columns
        and column != time_column
        and normalized_name(column) not in INDEX_COLUMNS
    ]
    y_values = {
        column: np.asarray([
            float(row[column]) if is_number(row.get(column)) else np.nan
            for row in source
        ])
        for column in numeric_columns
    }
    output: list[dict[str, Any]] = []
    for xx in q:
        record = {column: None for column in columns}
        for column in group_columns:
            record[column] = source[0].get(column)
        record[time_column] = float(xx)
        position = int(np.searchsorted(x, xx))
        low = max(0, position - 1)
        high = min(len(x) - 1, position)
        for column in numeric_columns:
            y = y_values[column]
            if np.isfinite(y[low]) and np.isfinite(y[high]):
                record[column] = float(np.interp(xx, x, y))
        output.append(record)
    return output


def group_columns_for(columns: list[str], rows: list[dict[str, Any]], time_column: str) -> list[str]:
    groups: list[str] = []
    for column in columns:
        low = normalized_name(column)
        if column == time_column or low in INDEX_COLUMNS:
            continue
        values = [row.get(column) for row in rows]
        if low in GROUP_COLUMNS or any(
            value is not None and not is_number(value) and not isinstance(value, bool)
            for value in values
        ):
            groups.append(column)
    return groups


def add_figure3_state_time(
    item_entry: dict[str, Any], dataset: dict[str, Any], rows: list[dict[str, Any]],
    csv_overrides: dict[str, Path] | None = None,
) -> tuple[list[str], list[dict[str, Any]]]:
    name = str(dataset["name"])
    if str(item_entry["item"]) != "Figure3" or not name.startswith("E_"):
        return [str(entry["name"]) for entry in dataset["columns"]], rows
    paired_name = name.replace("E_PPRtraces_states", "D_PPRtraces_").replace(
        "E_traces1_2__states", "D_traces1_2__"
    )
    paired = next(entry for entry in item_entry["datasets"] if entry["name"] == paired_name)
    _, paired_rows = read_csv(paired, csv_overrides)
    t0 = float(paired_rows[0]["time"]) / 1000.0
    augmented = [dict(row, time=t0 + index * 0.0002) for index, row in enumerate(rows)]
    return [*[str(entry["name"]) for entry in dataset["columns"]], "time"], augmented


def regular_block(
    item_entry: dict[str, Any],
    dataset: dict[str, Any],
    columns: list[str],
    rows: list[dict[str, Any]],
    expected: int,
) -> Block:
    item = str(item_entry["item"])
    name = str(dataset["name"])
    resample_bounds = (
        (100.0, 500.0)
        if item == "SupplementaryNote2" and name.endswith("poisson_zoom")
        else None
    )
    if item == "Figure1" and name.startswith("G_"):
        # These rectangular MATLAB exports pad the shorter demo curve with
        # empty rows to the longest stored vector. Only finite curve samples
        # belong to the plotted view (ending at 498.736 ms here).
        rows = [row for row in rows if row.get("column_1") is not None]
    if item == "ED7" and name.endswith("fluorescence_groundtruth") and len(rows) > expected:
        rows = [row for row in rows if row.get("observed_DeltaF_over_F") is not None]
    if item == "Figure2" and name == "I_PPR_inputs":
        selected = [row for row in rows if row.get("selected_for_I") is True]
        if len(selected) != expected:
            raise ValueError(f"Figure2 I selector produced {len(selected)} rows, expected {expected}")
        return Block(selected, "csv_exact_selected_for_I")
    if len(rows) == expected:
        return Block(rows, "csv_exact")
    time_columns = [column for column in columns if normalized_name(column) in TIME_COLUMNS]
    if len(time_columns) != 1:
        raise ValueError(
            f"{item}/{name}: needs an explicit source-to-view mapping "
            f"({len(rows)} CSV rows -> {expected} workbook rows; time={time_columns})"
        )
    time_column = time_columns[0]
    groups = group_columns_for(columns, rows, time_column)
    grouped: OrderedDict[tuple[Any, ...], list[dict[str, Any]]] = OrderedDict()
    for row in rows:
        grouped.setdefault(tuple(row.get(column) for column in groups), []).append(row)
    segments: list[tuple[list[dict[str, Any]], bool]] = []
    for group in grouped.values():
        value_columns = [
            column for column in columns
            if column != time_column
            and column not in groups
            and normalized_name(column) not in INDEX_COLUMNS
        ]
        if (
            (item == "Figure3" and name.startswith("Figure3_E_"))
            or (item == "ED3" and name.endswith("_states"))
        ):
            while group and all(group[0].get(column) is None for column in value_columns):
                group = group[1:]
            while group and all(group[-1].get(column) is None for column in value_columns):
                group = group[:-1]
        if "SEM_polygon" in name:
            middle = len(group) // 2
            segments.extend([(group[:middle], False), (group[middle:], True)])
        else:
            segments.append((group, False))
    output: list[dict[str, Any]] = []
    for segment, reverse in segments:
        per_segment = 200 if resample_bounds is not None else min(len(segment), 200)
        generated = resample_segment(
            segment,
            columns=columns,
            time_column=time_column,
            group_columns=groups,
            count=per_segment,
            bounds=resample_bounds,
        )
        output.extend(reversed(generated) if reverse else generated)
    if len(output) != expected:
        raise ValueError(
            f"{item}/{name}: declared time-series reduction produced {len(output)} rows; "
            f"workbook template expects {expected}"
        )
    return Block(output, "csv_resampled_uniform_time")


def location_block(
    item_entry: dict[str, Any],
    dataset: dict[str, Any],
    columns: list[str],
    rows: list[dict[str, Any]],
    location: dict[str, Any],
    target_headers: list[str],
) -> Block:
    expected = int(location["rows"])
    name = str(dataset["name"])
    normalized_targets = [normalized_name(header) for header in target_headers]
    if "source_path" in normalized_targets and not any(
        normalized_name(column) == "source_path" for column in columns
    ):
        rows = [dict(row, source_path=dataset.get("source")) for row in rows]
        columns = [*columns, "source_path"]
    if str(item_entry["item"]) == "SupplementaryNote2" and "cluster" in normalized_targets:
        cluster = {"B": 1, "C": 5, "D": 9}[name.split("_", 1)[0]]
        rows = [dict(row, cluster=cluster) for row in rows]
        columns = ["cluster", *columns]
    fluorescence = next(
        (header for header in target_headers if normalized_name(header).startswith("fluorescence_column")),
        None,
    )
    if fluorescence is not None:
        suffix = re.search(r"(\d+)$", normalized_name(fluorescence))
        if suffix is None:
            raise ValueError(f"Cannot identify calcium column in {fluorescence!r}")
        number = suffix.group(1)
        time_column = f"time_column{number}"
        fluorescence_column = f"fluorescence_column{number}"
        block_columns = [time_column, fluorescence_column]
        block_rows = [
            {time_column: row.get(time_column), fluorescence_column: row.get(fluorescence_column)}
            for row in rows
        ]
        result = resample_segment(
            block_rows,
            columns=block_columns,
            time_column=time_column,
            group_columns=[],
            count=expected,
        )
        for row in result:
            row["time"] = row[time_column]
        return Block(result, "csv_resampled_independent_time_column", f"{name}_calcium_column{number}")
    if len(dataset.get("workbook_locations", [])) > 1 and "avg_spike_count" in normalized_targets:
        return Block(rows, "csv_exact_posterior_bins")
    if len(dataset.get("workbook_locations", [])) > 1 and "dff" in normalized_targets:
        selected_columns = [column for column in columns if normalized_name(column) != "avg_spike_count"]
        return regular_block(item_entry, dataset, selected_columns, rows, expected)
    return regular_block(item_entry, dataset, columns, rows, expected)


def source_key(record: dict[str, Any], target_header: str, row_number: int) -> Any:
    key = normalized_name(target_header)
    if key == "interpolated_point_index":
        return row_number
    aliases = {"time": ("time", "time_s", "time_ms", "stored_time")}
    for column, value in record.items():
        if normalized_name(column) == key:
            return value
    for candidate in aliases.get(key, ()):
        for column, value in record.items():
            if normalized_name(column) == candidate:
                return value
    raise KeyError(f"No CSV-derived value for workbook column {target_header!r}")


def values_match(left: Any, right: Any) -> bool:
    def missing(value: Any) -> bool:
        return value is None or (
            isinstance(value, str) and value.strip().lower() in {"nan", "na", "n.a."}
        )

    if missing(left) and missing(right):
        return True
    if isinstance(left, str) and isinstance(right, bool):
        return left.strip().lower() == str(right).lower()
    if isinstance(right, str) and isinstance(left, bool):
        return right.strip().lower() == str(left).lower()
    def numeric(value: Any) -> float | None:
        if is_number(value):
            return float(value)
        if isinstance(value, str):
            try:
                return float(value.strip())
            except ValueError:
                return None
        return None

    left_number = numeric(left)
    right_number = numeric(right)
    if left_number is not None and right_number is not None:
        return math.isclose(left_number, right_number, rel_tol=1e-9, abs_tol=1e-11)
    return str(left) == str(right)


def representationally_same(left: Any, right: Any) -> bool:
    """Ignore only floating-point serialization drift, never substantive edits."""
    if left is None or right is None:
        return left is None and right is None
    if isinstance(left, bool) or isinstance(right, bool):
        return left == right
    if is_number(left) and is_number(right):
        return math.isclose(float(left), float(right), rel_tol=1e-15, abs_tol=0.0)
    return type(left) is type(right) and left == right


def presentation_value(existing: Any, derived: Any) -> Any:
    """Retain an intentional textual numeric format when numerically equivalent."""
    # Zero and IEEE-754 subnormal values are equivalent at the template
    # workbook's display precision. Preserve whichever representation the
    # publication layout uses while keeping the full-precision CSV untouched.
    if is_number(existing) and is_number(derived):
        if abs(float(existing)) < sys.float_info.min and abs(float(derived)) < sys.float_info.min:
            return existing
    if derived is None and isinstance(existing, str) and values_match(existing, derived):
        return existing
    if isinstance(existing, str) and isinstance(derived, bool) and values_match(existing, derived):
        return existing
    if isinstance(existing, str) and is_number(derived):
        try:
            existing_number = float(existing.strip())
        except ValueError:
            pass
        else:
            if representationally_same(existing_number, derived):
                return existing
    return derived


def block_cell_values(
    *,
    sheet: Any,
    start: int,
    target_headers: list[str],
    block: Block,
    dataset_name: str,
) -> tuple[dict[str, Any], list[str]]:
    values: dict[str, Any] = {}
    mismatches: list[str] = []
    label = block.dataset_label or dataset_name
    target_dataset_positions = [
        index for index, header in enumerate(target_headers, start=1)
        if normalized_name(header) == "dataset"
    ]
    source_dataset_count = max(
        (sum(normalized_name(column) == "dataset" for column in row) for row in block.rows),
        default=0,
    )
    pseudo_dataset_positions = set(
        target_dataset_positions[: max(0, len(target_dataset_positions) - source_dataset_count)]
    )

    def value_for(
        record: dict[str, Any], header: str, row_number: int, column_number: int
    ) -> Any:
        if column_number in pseudo_dataset_positions:
            return label
        value = source_key(record, header, row_number)
        return value

    for row_offset, record in enumerate(block.rows):
        for column_index, header in enumerate(target_headers, start=1):
            value = value_for(record, header, row_offset + 1, column_index)
            cell = sheet.cell(start + row_offset, column_index)
            values[cell.coordinate] = presentation_value(cell.value, value)
            if not values_match(cell.value, value):
                mismatches.append(
                    f"{dataset_name}: {sheet.title}!{cell.coordinate}: "
                    f"existing={cell.value!r}, CSV-derived={value!r}"
                )
    return values, mismatches


def _worksheet_parts(path: Path) -> dict[str, str]:
    """Map visible worksheet names to their OOXML ZIP members."""
    ns = {
        "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "p": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rels.findall("p:Relationship", ns)
    }
    result: dict[str, str] = {}
    for sheet in workbook.findall("m:sheets/m:sheet", ns):
        target = targets[sheet.attrib[f"{{{ns['r']}}}id"]].replace("\\", "/")
        if target.startswith("/"):
            member = target.lstrip("/")
        elif target.startswith("xl/"):
            member = target
        else:
            member = f"xl/{target}"
        result[sheet.attrib["name"]] = member
    return result


# Do not let a self-closing cell act as the opening tag for the next paired
# cell.  The negative lookahead is deliberately confined to the start tag.
CELL_RE = re.compile(
    rb"<(?P<tag>(?:[A-Za-z_][\w.-]*:)?c)\b(?![^>]*?/>)"
    rb"(?P<attrs>[^>]*\br=\"(?P<ref>[^\"]+)\"[^>]*)>"
    rb"(?P<body>.*?)</(?P=tag)>",
    re.S,
)
SELF_CLOSING_CELL_RE = re.compile(
    rb"<(?P<tag>(?:[A-Za-z_][\w.-]*:)?c)\b"
    rb"(?P<attrs>[^>]*\br=\"(?P<ref>[^\"]+)\"[^>]*)/>",
    re.S,
)
ROW_RE = re.compile(
    rb"<(?P<tag>(?:[A-Za-z_][\w.-]*:)?row)\b"
    rb"(?P<attrs>[^>]*\br=\"(?P<row>\d+)\"[^>]*)>"
    rb"(?P<body>.*?)</(?P=tag)>",
    re.S,
)
CELL_REFERENCE_RE = re.compile(r"^(?P<column>[A-Z]+)(?P<row>\d+)$")


def _serialize_cell(attrs: bytes, value: Any, tag: bytes = b"c") -> bytes:
    child_prefix = tag.rsplit(b":", 1)[0] + b":" if b":" in tag else b""
    attrs = re.sub(rb"\s+t=\"[^\"]*\"", b"", attrs)
    if value is None:
        return b"<" + tag + attrs + b"></" + tag + b">"
    if isinstance(value, bool):
        attrs += b' t="b"'
        body = b"<" + child_prefix + b"v>" + (b"1" if value else b"0") + b"</" + child_prefix + b"v>"
    elif is_number(value):
        number = float(value)
        if not math.isfinite(number):
            return b"<" + tag + attrs + b"></" + tag + b">"
        text = str(int(number)) if number.is_integer() and isinstance(value, int) else repr(number)
        body = b"<" + child_prefix + b"v>" + text.encode("utf-8") + b"</" + child_prefix + b"v>"
    else:
        attrs += b' t="inlineStr"'
        escaped = html.escape(str(value), quote=False)
        body = (
            b"<" + child_prefix + b"is><" + child_prefix
            + b't xml:space="preserve">' + escaped.encode("utf-8")
            + b"</" + child_prefix + b"t></" + child_prefix + b"is>"
        )
    return b"<" + tag + attrs + b">" + body + b"</" + tag + b">"


def _append_missing_rows(raw: bytes, references: Iterable[str]) -> bytes:
    """Append contiguous worksheet rows by cloning the final row's formatting.

    Source-data tables normally fit inside their workbook template.
    Figure 1's corrected panel-f statistics add two records, so its legend
    sheet needs two additional rows.  Only row/cell formatting is cloned; all
    values are cleared before the requested CSV-derived values are applied.
    """
    parsed = [CELL_REFERENCE_RE.fullmatch(reference) for reference in references]
    if any(match is None for match in parsed):
        raise ValueError("Cannot append cells with non-A1 references")
    target_rows = sorted({int(match.group("row")) for match in parsed if match is not None})
    existing = list(ROW_RE.finditer(raw))
    if not existing:
        raise ValueError("Cannot append rows to a worksheet without row templates")
    maximum = max(int(match.group("row")) for match in existing)
    missing_rows = [row for row in target_rows if row > maximum]
    if not missing_rows:
        return raw
    expected = list(range(maximum + 1, max(missing_rows) + 1))
    if missing_rows != expected:
        raise ValueError(
            f"Worksheet extension must be contiguous after row {maximum}: {missing_rows}"
        )
    template = max(existing, key=lambda match: int(match.group("row")))

    def blank_cell(match: re.Match[bytes], row_number: int) -> bytes:
        attrs = re.sub(
            rb'\br=\"[A-Z]+\d+\"',
            lambda ref: b'r="' + re.match(rb"[A-Z]+", ref.group(0).split(b'"')[1]).group(0)
            + str(row_number).encode("ascii") + b'"',
            match.group("attrs"),
            count=1,
        )
        return _serialize_cell(attrs, None, match.group("tag"))

    appended: list[bytes] = []
    for row_number in expected:
        attrs = re.sub(
            rb'\br=\"\d+\"',
            f'r="{row_number}"'.encode("ascii"),
            template.group("attrs"),
            count=1,
        )
        body = CELL_RE.sub(lambda match: blank_cell(match, row_number), template.group("body"))
        body = SELF_CLOSING_CELL_RE.sub(
            lambda match: blank_cell(match, row_number), body
        )
        appended.append(
            b"<" + template.group("tag") + attrs + b">" + body
            + b"</" + template.group("tag") + b">"
        )
    closing = re.search(rb"</(?:[A-Za-z_][\w.-]*:)?sheetData>", raw)
    if closing is None:
        raise ValueError("Cannot find worksheet sheetData closing tag")
    raw = raw[:closing.start()] + b"".join(appended) + raw[closing.start():]

    def extend_dimension(match: re.Match[bytes]) -> bytes:
        start, end_column, end_row = match.groups()
        return start + end_column + str(expected[-1]).encode("ascii") + b'"'

    raw, count = re.subn(
        rb'(<dimension\b[^>]*\bref="[A-Z]+\d+:)([A-Z]+)(\d+)"',
        extend_dimension,
        raw,
        count=1,
    )
    if count != 1:
        raise ValueError("Cannot extend worksheet dimension")
    return raw


def _patch_sheet_xml(raw: bytes, changes: dict[str, Any]) -> bytes:
    """Surgically replace only requested cell elements in a worksheet part."""
    remaining = dict(changes)

    existing_references = {
        match.group("ref").decode("ascii") for match in CELL_RE.finditer(raw)
    } | {
        match.group("ref").decode("ascii")
        for match in SELF_CLOSING_CELL_RE.finditer(raw)
    }
    absent = [reference for reference in remaining if reference not in existing_references]
    if absent:
        raw = _append_missing_rows(raw, absent)

    def replace(match: re.Match[bytes]) -> bytes:
        ref = match.group("ref").decode("ascii")
        if ref not in remaining:
            return match.group(0)
        value = remaining.pop(ref)
        return _serialize_cell(match.group("attrs"), value, match.group("tag"))

    raw = CELL_RE.sub(replace, raw)
    raw = SELF_CLOSING_CELL_RE.sub(replace, raw)
    if remaining:
        missing = ", ".join(sorted(remaining)[:8])
        raise ValueError(f"Cannot update absent worksheet cells: {missing}")
    return raw


def write_xlsx_preserving_package(
    template: Path, destination: Path, changes_by_sheet: dict[str, dict[str, Any]]
) -> None:
    """Write changed cell values while copying every other ZIP member verbatim."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    if not any(changes_by_sheet.values()):
        shutil.copy2(template, destination)
        return
    parts = _worksheet_parts(template)
    changes_by_part = {
        parts[sheet]: changes for sheet, changes in changes_by_sheet.items() if changes
    }
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx", dir=destination.parent) as stream:
        temporary = Path(stream.name)
    try:
        with zipfile.ZipFile(template, "r") as source, zipfile.ZipFile(temporary, "w") as target:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename in changes_by_part:
                    try:
                        payload = _patch_sheet_xml(payload, changes_by_part[info.filename])
                    except ValueError as exc:
                        raise ValueError(f"{info.filename}: {exc}") from exc
                target.writestr(info, payload)
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)


def build_one(
    item_entry: dict[str, Any], output_dir: Path, *,
    csv_overrides: dict[str, Path] | None = None,
) -> dict[str, Any]:
    item = ITEM_BY_SOURCE_DATA[str(item_entry["item"])]
    item_dir = REPO_ROOT / str(item["path"])
    template = item_dir / "source_data" / str(item_entry["workbook"])
    destination = output_dir / str(item_entry["workbook"])
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(template, data_only=False)
    changes_by_sheet: dict[str, dict[str, Any]] = {}
    all_mismatches: list[str] = []
    dataset_reports: list[dict[str, Any]] = []
    for dataset in item_entry["datasets"]:
        locations = dataset.get("workbook_locations", [])
        if not locations:
            continue
        columns, source_rows = read_csv(dataset, csv_overrides)
        columns, source_rows = add_figure3_state_time(
            item_entry, dataset, source_rows, csv_overrides
        )
        mismatch_start = len(all_mismatches)
        modes: list[str] = []
        rebuilt_rows = 0
        for location in locations:
            sheet = resolve_sheet(workbook, str(location["sheet"]))
            target_headers = [
                str(sheet.cell(int(location["header_row"]), column).value)
                for column in range(1, int(location["columns"]) + 1)
            ]
            block = location_block(
                item_entry, dataset, columns, source_rows, location, target_headers
            )
            if len(block.rows) != int(location["rows"]):
                raise ValueError(
                    f"{item_entry['item']}/{dataset['name']}: generated {len(block.rows)} "
                    f"rows for {location['sheet']}, expected {location['rows']}"
                )
            cell_values, mismatches = block_cell_values(
                sheet=sheet,
                start=int(location["data_start_row"]),
                target_headers=target_headers,
                block=block,
                dataset_name=str(dataset["name"]),
            )
            all_mismatches.extend(mismatches)
            sheet_changes = changes_by_sheet.setdefault(sheet.title, {})
            for coordinate, value in cell_values.items():
                if not representationally_same(sheet[coordinate].value, value):
                    sheet_changes[coordinate] = value
            modes.append(block.mode)
            rebuilt_rows += len(block.rows)
        dataset_reports.append({
            "name": dataset["name"],
            "csv": dataset["csv"],
            "mode": "+".join(OrderedDict.fromkeys(modes)),
            "workbook_rows": rebuilt_rows,
            "template_value_differences_before_write": len(all_mismatches) - mismatch_start,
        })
    workbook.close()
    write_xlsx_preserving_package(template, destination, changes_by_sheet)
    changed_cells = sum(len(changes) for changes in changes_by_sheet.values())
    byte_identical = sha256(destination) == sha256(template)
    if changed_cells == 0 and not byte_identical:
        raise ValueError(f"Unchanged workbook package is not byte-identical: {destination}")
    return {
        "item": item_entry["item"],
        "publication_id": item["publication_id"],
        "output": report_path(destination),
        "template": report_path(template),
        "status": "regenerated_from_csv",
        "changed_cells": changed_cells,
        "sha256": sha256(destination),
        "datasets": dataset_reports,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--item", action="append", help="Source Data item ID; may be repeated")
    parser.add_argument("--output-dir", type=Path, default=REPO_ROOT / "build" / "source_data")
    parser.add_argument("--report", type=Path, help="Optional JSON build summary path")
    args = parser.parse_args()
    registry = load_registry()
    selected = set(args.item or [])
    unknown = sorted(selected - {str(item["item"]) for item in registry})
    if unknown:
        parser.error("Unknown Source Data item IDs: " + ", ".join(unknown))
    reports = []
    for item in registry:
        if selected and item["item"] not in selected:
            continue
        try:
            reports.append(build_one(item, args.output_dir))
        except Exception as exc:
            raise RuntimeError(f"Workbook build failed for {item['item']}: {exc}") from exc
    if args.report is not None:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(reports, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({report["publication_id"]: report["status"] for report in reports}, indent=2))


if __name__ == "__main__":
    main()
